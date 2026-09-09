"""
loaders.py — parses every supported file type and attaches metadata
(filename, source_type, reliability_tier) in one place.

This replaces parsers.py and reliability_tags.py entirely — keeping
parsing logic and reliability tagging together avoids two files
disagreeing about the same document.
"""

import os
from pypdf import PdfReader
from pdf2image import convert_from_path
from docx import Document
from PIL import Image
import pytesseract

# Windows doesn't add these to PATH automatically during install.
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
POPPLER_PATH = r"C:\Users\ranju\Desktop\Poppler\poppler-26.07.0\Library\bin"


def load_pdf(filepath):
    """Real text extraction, with OCR fallback for scanned pages."""
    results = []
    reader = PdfReader(filepath)

    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        text = text.strip()

        if len(text) < 20:  # basically empty -> likely a scanned page
            images = convert_from_path(
                filepath, first_page=i + 1, last_page=i + 1,
                poppler_path=POPPLER_PATH
            )
            text = pytesseract.image_to_string(images[0]).strip()

        results.append({"text": text, "page_number": i + 1})

    return results


def load_docx(filepath):
    """DOCX has no native page concept, so page_number is None."""
    doc = Document(filepath)
    parts = []

    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                parts.append(row_text)

    full_text = "\n".join(parts)
    return [{"text": full_text, "page_number": None}]


def load_text(filepath):
    """TXT and MD files — read as-is, keep markdown syntax (headers carry meaning)."""
    with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read().strip()
    return [{"text": text, "page_number": None}]


def load_image(filepath):
    """Standalone image OCR (not inside a PDF)."""
    img = Image.open(filepath)
    text = pytesseract.image_to_string(img).strip()
    return [{"text": text, "page_number": None}]


def get_reliability_tier(source_type):
    """source_type = top-level folder name (codex, wiki, chronicles, ephemera, images)."""
    return {
        "codex": "official",
        "wiki": "reference",
        "chronicles": "narrative",
        "ephemera": "unreliable",
        "images": "unreliable",  # no author context on standalone plates — treat cautiously
    }.get(source_type, "unknown")


def load_document(filepath, source_type):
    """
    The router. Returns a list of chunk-ready dicts:
    [{"text", "filename", "source_type", "reliability_tier", "page_number"}, ...]
    """
    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".pdf":
        pages = load_pdf(filepath)
    elif ext == ".docx":
        pages = load_docx(filepath)
    elif ext in (".txt", ".md"):
        pages = load_text(filepath)
    elif ext in (".jpg", ".jpeg", ".png"):
        pages = load_image(filepath)
    else:
        print(f"Skipping unsupported file: {filepath}")
        return []

    reliability_tier = get_reliability_tier(source_type)
    docs = []
    for page in pages:
        if page["text"]:  # skip empty extractions
            docs.append({
                "text": page["text"],
                "filename": os.path.basename(filepath),
                "source_type": source_type,
                "reliability_tier": reliability_tier,
                "page_number": page["page_number"],
            })
    return docs